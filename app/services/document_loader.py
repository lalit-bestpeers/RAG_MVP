import json
import logging
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
import pytesseract

from app.core.config import settings

logger = logging.getLogger("rag.loader")

try:
    import pymupdf
except ImportError:  # pragma: no cover - older pymupdf versions
    import fitz as pymupdf


@dataclass
class ExtractedPage:
    page_number: int
    text: str


def extract_pages(file_path: Path) -> list[ExtractedPage]:
    ext = file_path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext in {".txt", ".md", ".html"}:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        return [ExtractedPage(page_number=1, text=text)]
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext == ".pptx":
        return _extract_pptx(file_path)
    if ext in {".png", ".jpg", ".jpeg"}:
        return [ExtractedPage(page_number=1, text=_ocr_image(file_path))]

    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_path: Path) -> list[ExtractedPage]:
    if settings.document_parser.lower() == "opendataloader":
        try:
            pages = _extract_with_opendataloader(file_path)
            if any(p.text.strip() for p in pages):
                return pages
            logger.info("OpenDataLoader produced no text; falling back to PyMuPDF/OCR")
        except Exception as exc:
            logger.warning("OpenDataLoader extraction failed (%s); falling back to PyMuPDF/OCR", exc)
    return _extract_with_pymupdf(file_path, ocr=True)


def _extract_with_opendataloader(file_path: Path) -> list[ExtractedPage]:
    import opendataloader_pdf

    with tempfile.TemporaryDirectory() as tmp:
        opendataloader_pdf.convert(
            input_path=[str(file_path)],
            output_dir=tmp,
            format="json",
            quiet=True,
            image_output="off",
        )
        json_files = list(Path(tmp).glob("*.json"))
        if not json_files:
            raise RuntimeError("OpenDataLoader produced no JSON output")

        pages_by_number: dict[int, list[str]] = defaultdict(list)
        total_pages = 0

        for json_file in json_files:
            with open(json_file, encoding="utf-8") as fh:
                data = json.load(fh)
            documents = data if isinstance(data, list) else [data]
            for doc in documents:
                if not isinstance(doc, dict):
                    continue
                total_pages = max(total_pages, int(doc.get("number of pages") or 0))
                _walk_elements(doc.get("kids") or [], pages_by_number)

        if not pages_by_number:
            raise RuntimeError("OpenDataLoader extracted no text")

        pages = []
        for page in range(1, total_pages + 1):
            text = "\n".join(pages_by_number.get(page, [])).strip()
            pages.append(ExtractedPage(page_number=page, text=text))
        return [p for p in pages if p.text] or pages


def _walk_elements(elements, pages_by_number: dict[int, list[str]]) -> None:
    for element in elements:
        if not isinstance(element, dict):
            continue
        content = element.get("content")
        page = element.get("page number")
        if isinstance(content, str) and content.strip() and isinstance(page, int):
            pages_by_number[page].append(content.strip())
        for key in ("kids", "list items", "rows", "cells"):
            children = element.get(key)
            if isinstance(children, list):
                _walk_elements(children, pages_by_number)


def _extract_with_pymupdf(file_path: Path, ocr: bool = True) -> list[ExtractedPage]:
    pdf = pymupdf.open(str(file_path))
    pages = []
    for page in pdf:
        text = page.get_text("text").strip()
        if ocr and len(text) < 20:
            try:
                text = _ocr_pdf_page(page)
            except Exception as exc:
                logger.warning("OCR failed for a page: %s", exc)
        pages.append(ExtractedPage(page_number=page.number + 1, text=text))
    pdf.close()
    return pages


def _ocr_pdf_page(page) -> str:
    pix = page.get_pixmap(
        matrix=pymupdf.Matrix(300 / 72, 300 / 72), colorspace=pymupdf.csRGB, alpha=False
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return pytesseract.image_to_string(image)


def _ocr_image(file_path: Path) -> str:
    image = Image.open(str(file_path))
    return pytesseract.image_to_string(image)


def _extract_docx(file_path: Path) -> list[ExtractedPage]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(file_path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [ExtractedPage(page_number=1, text=text)]


def _extract_pptx(file_path: Path) -> list[ExtractedPage]:
    from pptx import Presentation

    prs = Presentation(str(file_path))
    pages = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
        pages.append(ExtractedPage(page_number=idx, text="\n".join(parts)))
    return pages
